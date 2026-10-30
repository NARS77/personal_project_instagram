#!/usr/bin/env python3
"""
MeatNDoor Lead Outreach Manager.

This module provides end-to-end lead ingestion, location-based campaign
eligibility, personalized message drafting, review queue management,
Do-Not-Contact (DNC) enforcement, authorized Meta messaging integration,
and dual synchronization (Google Sheets & Excel).

Guiding Principles:
1. Incremental and idempotent ingestion from existing checkpoint databases.
2. Reuses cached location-checking intelligence from mumbai_checks.
3. Strict campaign eligibility:
   - Likely Mumbai: eligible for Mumbai campaign review.
   - Serves Mumbai: review separately (business service != proof of residence).
   - Possibly Mumbai: lower-confidence manual review.
   - Not enough evidence: location unknown.
   - Likely outside Mumbai: suppressed from Mumbai campaigns.
   - Do Not Contact: suppressed from all campaigns.
4. No automated spam, no bulk follow bots, no unauthorized private API scraping.
5. SQLite is the persistent source of truth; zero CSV exports.
"""

import json
import logging
import os
import re
import sqlite3
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional, Set, Tuple

logger = logging.getLogger("instagram_comments.lead_outreach")

# Canonical Paths
DEFAULT_SPREADSHEET_ID = "1sLCfz9Eho4OWjsDTILIVKNRtwVYbCoibFK6z6un7c_Y"
DEFAULT_EXCEL_PATH = "exports/instagram_research.xlsx"

# Standard Messaging Templates
MUMBAI_OUTREACH_TEMPLATE = (
    "Hi! I'm reaching out from MeatNDoor. We offer a range of meat products. "
    "If Mumbai is your area, I can share our current product list and help you "
    "check availability for your location. Would you be interested in the details?"
)

GENERIC_OUTREACH_TEMPLATE = (
    "Hi! I'm reaching out from MeatNDoor. We offer a range of meat products. "
    "If you're interested, I can share our current product list and help you "
    "check availability in your area. Would you like to see the details?"
)

# Campaigns
CAMPAIGN_MUMBAI = "MeatNDoor Mumbai Launch"
CAMPAIGN_GENERIC = "MeatNDoor General Outreach"

# Eligibility Constants
ELIGIBILITY_MUMBAI_TARGET = "Eligible (Mumbai Target)"
ELIGIBILITY_BUSINESS_REVIEW = "Review (Business/Serves Mumbai)"
ELIGIBILITY_LOW_CONF_REVIEW = "Review (Low Confidence)"
ELIGIBILITY_PENDING_LOCATION = "Pending Location Verification"
ELIGIBILITY_SUPPRESSED_OUTSIDE = "Suppressed (Outside Mumbai)"
ELIGIBILITY_SUPPRESSED_DNC = "Suppressed (Do Not Contact)"

# Review Status Constants
STATUS_PENDING = "Pending Review"
STATUS_APPROVED = "Approved"
STATUS_REJECTED = "Rejected"
STATUS_CONTACTED = "Contacted"
STATUS_REPLIED = "Replied"
STATUS_INTERESTED = "Interested"
STATUS_NOT_INTERESTED = "Not Interested"
STATUS_DNC = "Do Not Contact"

ALL_REVIEW_STATUSES = [
    STATUS_PENDING,
    STATUS_APPROVED,
    STATUS_REJECTED,
    STATUS_CONTACTED,
    STATUS_REPLIED,
    STATUS_INTERESTED,
    STATUS_NOT_INTERESTED,
    STATUS_DNC,
]

# Messaging Workflow Status
MSG_STATUS_MANUAL = "Manual / Opt-In Required"
MSG_STATUS_READY = "Authorized API Ready"
MSG_STATUS_SENT = "Sent via API"
MSG_STATUS_FAILED = "Failed"


# Custom Exceptions
class LeadOutreachError(Exception):
    """Base exception for Lead Outreach Manager."""
    pass


class DNCSuppressionError(LeadOutreachError):
    """Raised when an action is attempted on an account suppressed by Do Not Contact."""
    pass


class CampaignSuppressionError(LeadOutreachError):
    """Raised when attempting to contact an ineligible or outside-region account."""
    pass


class DuplicateSendBlockedError(LeadOutreachError):
    """Raised to prevent duplicate contact attempts to the same account."""
    pass


class UnapprovedDraftError(LeadOutreachError):
    """Raised when attempting outreach before the draft has been reviewed and approved."""
    pass


class MetaAuthorizationError(LeadOutreachError):
    """Raised when required Meta Graph API permissions or credentials are missing."""
    pass


def get_canonical_checkpoints_dir() -> Path:
    """Finds or initializes the canonical .checkpoints directory."""
    candidates = [
        Path(__file__).resolve().parent / ".checkpoints",
        Path.cwd() / "api-anything" / ".checkpoints",
        Path.cwd() / ".checkpoints",
    ]
    for c in candidates:
        if c.exists() and c.is_dir():
            return c.resolve()
    target = Path(__file__).resolve().parent / ".checkpoints"
    target.mkdir(parents=True, exist_ok=True)
    return target.resolve()


def get_canonical_outreach_db_path() -> Path:
    """Returns canonical path to .checkpoints/lead_outreach.db."""
    cp_dir = get_canonical_checkpoints_dir()
    return cp_dir / "lead_outreach.db"


def get_canonical_mumbai_db_path() -> Path:
    """Finds canonical mumbai_check_checkpoint.db path."""
    cp_dir = get_canonical_checkpoints_dir()
    db_file = cp_dir / "mumbai_check_checkpoint.db"
    if db_file.exists():
        return db_file.resolve()
    # Fallback to parent or current
    for candidate in [
        Path.cwd() / "api-anything" / ".checkpoints" / "mumbai_check_checkpoint.db",
        Path.cwd() / ".checkpoints" / "mumbai_check_checkpoint.db",
    ]:
        if candidate.exists():
            return candidate.resolve()
    return db_file


def normalize_username(username: str) -> str:
    """Normalizes an Instagram username for consistent indexing."""
    if not username:
        return ""
    clean = username.strip().lstrip("@").lower()
    return re.sub(r"[^a-z0-9._]", "", clean)


class LeadOutreachDB:
    """
    SQLite-backed repository for MeatNDoor Lead Outreach.
    Manages lead records, DNC registry, and operator audit trail.
    """

    def __init__(self, db_path: Optional[Path] = None):
        self.db_path = Path(db_path) if db_path else get_canonical_outreach_db_path()
        self.db_path.parent.mkdir(parents=True, exist_ok=True)
        self._init_db()

    def _get_connection(self) -> sqlite3.Connection:
        conn = sqlite3.connect(str(self.db_path), timeout=30.0)
        conn.row_factory = sqlite3.Row
        conn.execute("PRAGMA journal_mode = WAL")
        conn.execute("PRAGMA synchronous = NORMAL")
        conn.execute("PRAGMA foreign_keys = ON")
        return conn

    def _init_db(self) -> None:
        """Initializes tables for leads, DNC registry, and operator audit logs."""
        with self._get_connection() as conn:
            conn.executescript("""
                CREATE TABLE IF NOT EXISTS outreach_leads (
                    normalized_username TEXT PRIMARY KEY,
                    username TEXT NOT NULL,
                    profile_url TEXT NOT NULL,
                    direct_message_url TEXT NOT NULL,
                    source_brands TEXT NOT NULL,
                    source_posts TEXT NOT NULL,
                    total_comments INTEGER DEFAULT 1,
                    comment_types TEXT DEFAULT '',
                    sample_comment_text TEXT DEFAULT '',
                    mumbai_classification TEXT NOT NULL DEFAULT 'Not enough evidence',
                    mumbai_confidence TEXT NOT NULL DEFAULT 'Low',
                    stated_location TEXT DEFAULT '',
                    evidence_summary TEXT DEFAULT '',
                    campaign TEXT NOT NULL DEFAULT 'MeatNDoor Mumbai Launch',
                    campaign_eligibility TEXT NOT NULL,
                    drafted_message TEXT NOT NULL,
                    review_status TEXT NOT NULL DEFAULT 'Pending Review',
                    is_dnc INTEGER DEFAULT 0,
                    dnc_reason TEXT DEFAULT '',
                    messaging_status TEXT NOT NULL DEFAULT 'Manual / Opt-In Required',
                    last_action TEXT DEFAULT 'Ingested',
                    last_action_at TEXT NOT NULL,
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS do_not_contact (
                    normalized_username TEXT PRIMARY KEY,
                    username TEXT NOT NULL,
                    reason TEXT NOT NULL,
                    added_at TEXT NOT NULL,
                    added_by TEXT DEFAULT 'Operator'
                );

                CREATE TABLE IF NOT EXISTS outreach_audit_log (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    normalized_username TEXT NOT NULL,
                    action TEXT NOT NULL,
                    previous_status TEXT,
                    new_status TEXT,
                    notes TEXT,
                    operator TEXT DEFAULT 'Operator',
                    created_at TEXT NOT NULL
                );

                CREATE INDEX IF NOT EXISTS idx_leads_campaign_eligibility ON outreach_leads(campaign_eligibility);
                CREATE INDEX IF NOT EXISTS idx_leads_review_status ON outreach_leads(review_status);
                CREATE INDEX IF NOT EXISTS idx_leads_is_dnc ON outreach_leads(is_dnc);
                CREATE INDEX IF NOT EXISTS idx_audit_username ON outreach_audit_log(normalized_username);
            """)

    # -------------------------------------------------------------------------
    # Ingestion & Cross-Reference Engine
    # -------------------------------------------------------------------------
    def ingest_from_checkpoints(
        self,
        checkpoint_dir: Optional[Path] = None,
        source_accounts: Optional[List[str]] = None,
        mumbai_db_path: Optional[Path] = None,
    ) -> Dict[str, int]:
        """
        Reads commenters and reply authors from existing checkpoint DBs.
        Deduplicates usernames, links source posts/brands, joins with cached
        Mumbai location evaluations, applies eligibility rules, and drafts messages.

        Returns ingestion statistics: {
            "total_raw_found": int,
            "leads_ingested_or_updated": int,
            "new_leads_added": int,
            "dnc_suppressed": int,
            "mumbai_eligible": int
        }
        """
        cp_dir = Path(checkpoint_dir) if checkpoint_dir else get_canonical_checkpoints_dir()
        m_db_path = Path(mumbai_db_path) if mumbai_db_path else get_canonical_mumbai_db_path()

        # Load cached Mumbai checks
        mumbai_cache: Dict[str, Dict[str, Any]] = {}
        if m_db_path.exists():
            try:
                with sqlite3.connect(str(m_db_path), timeout=15.0) as mconn:
                    mconn.row_factory = sqlite3.Row
                    rows = mconn.execute(
                        "SELECT normalized_username, checked_username, classification, confidence, "
                        "stated_location, evidence_summary, has_error FROM mumbai_checks"
                    ).fetchall()
                    for r in rows:
                        norm = r["normalized_username"] or normalize_username(r["checked_username"])
                        mumbai_cache[norm] = dict(r)
            except Exception as e:
                logger.warning("Could not load Mumbai checks from %s: %s", m_db_path, e)

        # Load DNC set
        dnc_map = self.get_dnc_map()

        # Identify checkpoint DBs to process
        db_files = list(cp_dir.glob("*_checkpoint.db"))
        target_files = []
        for f in db_files:
            if f.name in ("mumbai_check_checkpoint.db", "lead_outreach.db"):
                continue
            stem = f.stem.replace("_checkpoint", "")
            if source_accounts:
                norm_sources = {normalize_username(s) for s in source_accounts}
                if normalize_username(stem) not in norm_sources:
                    continue
            target_files.append((f, stem))

        # Ingest commenters
        raw_found = 0
        leads_agg: Dict[str, Dict[str, Any]] = {}

        for db_file, brand_name in target_files:
            try:
                with sqlite3.connect(str(db_file), timeout=15.0) as conn:
                    conn.row_factory = sqlite3.Row
                    # Check if 'comments' table exists
                    table_check = conn.execute(
                        "SELECT name FROM sqlite_master WHERE type='table' AND name='comments'"
                    ).fetchone()
                    if not table_check:
                        continue

                    rows = conn.execute(
                        "SELECT username, shortcode, text, record_type FROM comments"
                    ).fetchall()

                    for r in rows:
                        u = r["username"]
                        if not u:
                            continue
                        norm_u = normalize_username(u)
                        # Rule: Do not treat the source brand account as a lead
                        if norm_u == normalize_username(brand_name):
                            continue

                        raw_found += 1
                        shortcode = r["shortcode"] or ""
                        rec_type = (r["record_type"] or "comment").strip().lower()
                        comment_text = (r["text"] or "").strip()

                        if norm_u not in leads_agg:
                            leads_agg[norm_u] = {
                                "username": u.strip().lstrip("@"),
                                "normalized_username": norm_u,
                                "source_brands": set(),
                                "source_posts": set(),
                                "total_comments": 0,
                                "comment_types": set(),
                                "sample_comment_text": comment_text,
                            }

                        lead = leads_agg[norm_u]
                        lead["source_brands"].add(brand_name)
                        if shortcode:
                            lead["source_posts"].add(shortcode)
                        lead["total_comments"] += 1
                        lead["comment_types"].add(rec_type)
                        if comment_text and not lead["sample_comment_text"]:
                            lead["sample_comment_text"] = comment_text
            except Exception as e:
                logger.error("Error ingesting from checkpoint DB %s: %s", db_file, e)

        # Persist and classify leads
        now_utc = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")
        new_count = 0
        updated_count = 0
        dnc_count = 0
        mumbai_eligible_count = 0

        with self._get_connection() as conn:
            for norm_u, data in leads_agg.items():
                existing = conn.execute(
                    "SELECT * FROM outreach_leads WHERE normalized_username = ?", (norm_u,)
                ).fetchone()

                # Location Intelligence from cache
                m_info = mumbai_cache.get(norm_u, {})
                classification = m_info.get("classification", "Not enough evidence")
                confidence = m_info.get("confidence", "Low")
                stated_location = m_info.get("stated_location", "")
                evidence_summary = m_info.get("evidence_summary", "")
                has_error = bool(m_info.get("has_error", 0))

                # Safety rule: Never convert transient errors into confirmed outside Mumbai
                if has_error and classification == "Likely outside Mumbai":
                    classification = "Not enough evidence"
                    confidence = "Low"

                # Check DNC
                is_dnc = 1 if norm_u in dnc_map else 0
                dnc_reason = dnc_map.get(norm_u, "")

                # Determine Campaign Eligibility
                if is_dnc:
                    eligibility = ELIGIBILITY_SUPPRESSED_DNC
                    dnc_count += 1
                elif classification == "Likely Mumbai":
                    eligibility = ELIGIBILITY_MUMBAI_TARGET
                    mumbai_eligible_count += 1
                elif classification == "Serves Mumbai":
                    eligibility = ELIGIBILITY_BUSINESS_REVIEW
                elif classification == "Possibly Mumbai":
                    eligibility = ELIGIBILITY_LOW_CONF_REVIEW
                elif classification == "Likely outside Mumbai":
                    eligibility = ELIGIBILITY_SUPPRESSED_OUTSIDE
                else:
                    eligibility = ELIGIBILITY_PENDING_LOCATION

                # Determine Message Draft
                if eligibility in (ELIGIBILITY_MUMBAI_TARGET, ELIGIBILITY_BUSINESS_REVIEW, ELIGIBILITY_LOW_CONF_REVIEW):
                    draft = MUMBAI_OUTREACH_TEMPLATE
                else:
                    draft = GENERIC_OUTREACH_TEMPLATE

                brands_json = json.dumps(sorted(list(data["source_brands"])))
                posts_json = json.dumps(sorted(list(data["source_posts"])))
                comment_types_str = ", ".join(sorted(list(data["comment_types"])))
                profile_url = f"https://www.instagram.com/{data['username']}/"
                dm_url = f"https://ig.me/m/{data['username']}"

                if existing:
                    # Idempotent update: merge brands and posts, update comments, preserve existing review state and custom draft!
                    exist_brands = set(json.loads(existing["source_brands"] or "[]"))
                    exist_posts = set(json.loads(existing["source_posts"] or "[]"))
                    merged_brands = json.dumps(sorted(list(exist_brands | data["source_brands"])))
                    merged_posts = json.dumps(sorted(list(exist_posts | data["source_posts"])))
                    tot_comments = max(existing["total_comments"], data["total_comments"])

                    # If operator has edited draft or progressed status, preserve it
                    cur_status = existing["review_status"]
                    if is_dnc:
                        cur_status = STATUS_DNC
                    cur_draft = existing["drafted_message"] if existing["drafted_message"] else draft

                    conn.execute("""
                        UPDATE outreach_leads SET
                            username = ?,
                            profile_url = ?,
                            direct_message_url = ?,
                            source_brands = ?,
                            source_posts = ?,
                            total_comments = ?,
                            comment_types = ?,
                            mumbai_classification = ?,
                            mumbai_confidence = ?,
                            stated_location = ?,
                            evidence_summary = ?,
                            campaign_eligibility = ?,
                            drafted_message = ?,
                            review_status = ?,
                            is_dnc = ?,
                            dnc_reason = ?,
                            updated_at = ?
                        WHERE normalized_username = ?
                    """, (
                        data["username"], profile_url, dm_url, merged_brands, merged_posts,
                        tot_comments, comment_types_str, classification, confidence,
                        stated_location, evidence_summary, eligibility, cur_draft,
                        cur_status, is_dnc, dnc_reason, now_utc, norm_u
                    ))
                    updated_count += 1
                else:
                    initial_status = STATUS_DNC if is_dnc else STATUS_PENDING
                    conn.execute("""
                        INSERT INTO outreach_leads (
                            normalized_username, username, profile_url, direct_message_url,
                            source_brands, source_posts, total_comments, comment_types,
                            sample_comment_text, mumbai_classification, mumbai_confidence,
                            stated_location, evidence_summary, campaign, campaign_eligibility,
                            drafted_message, review_status, is_dnc, dnc_reason,
                            messaging_status, last_action, last_action_at, created_at, updated_at
                        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """, (
                        norm_u, data["username"], profile_url, dm_url,
                        brands_json, posts_json, data["total_comments"], comment_types_str,
                        data["sample_comment_text"], classification, confidence,
                        stated_location, evidence_summary, CAMPAIGN_MUMBAI, eligibility,
                        draft, initial_status, is_dnc, dnc_reason,
                        MSG_STATUS_MANUAL, "Ingested", now_utc, now_utc, now_utc
                    ))
                    new_count += 1

        return {
            "total_raw_found": raw_found,
            "leads_ingested_or_updated": len(leads_agg),
            "new_leads_added": new_count,
            "dnc_suppressed": dnc_count,
            "mumbai_eligible": mumbai_eligible_count,
        }

    # -------------------------------------------------------------------------
    # Do Not Contact (DNC) Registry Management
    # -------------------------------------------------------------------------
    def get_dnc_map(self) -> Dict[str, str]:
        """Returns map of normalized_username -> reason for all DNC records."""
        with self._get_connection() as conn:
            rows = conn.execute("SELECT normalized_username, reason FROM do_not_contact").fetchall()
            return {r["normalized_username"]: r["reason"] for r in rows}

    def add_to_dnc(
        self, username: str, reason: str = "User requested Do Not Contact", operator: str = "Operator"
    ) -> bool:
        """Adds account to persistent DNC registry and updates lead status."""
        norm = normalize_username(username)
        if not norm:
            return False

        now_utc = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")
        with self._get_connection() as conn:
            conn.execute("""
                INSERT INTO do_not_contact (normalized_username, username, reason, added_at, added_by)
                VALUES (?, ?, ?, ?, ?)
                ON CONFLICT(normalized_username) DO UPDATE SET
                    reason = excluded.reason,
                    added_at = excluded.added_at,
                    added_by = excluded.added_by
            """, (norm, username.strip().lstrip("@"), reason, now_utc, operator))

            # Update existing lead if present
            lead = conn.execute("SELECT review_status FROM outreach_leads WHERE normalized_username = ?", (norm,)).fetchone()
            prev_status = lead["review_status"] if lead else None

            conn.execute("""
                UPDATE outreach_leads SET
                    is_dnc = 1,
                    dnc_reason = ?,
                    campaign_eligibility = ?,
                    review_status = ?,
                    last_action = 'Added to DNC',
                    last_action_at = ?,
                    updated_at = ?
                WHERE normalized_username = ?
            """, (reason, ELIGIBILITY_SUPPRESSED_DNC, STATUS_DNC, now_utc, now_utc, norm))

            conn.execute("""
                INSERT INTO outreach_audit_log (
                    normalized_username, action, previous_status, new_status, notes, operator, created_at
                ) VALUES (?, 'ADD_DNC', ?, ?, ?, ?, ?)
            """, (norm, prev_status, STATUS_DNC, reason, operator, now_utc))

        logger.info("Added @%s to Do Not Contact (reason: %s)", username, reason)
        return True

    def remove_from_dnc(self, username: str, operator: str = "Operator") -> bool:
        """Removes account from DNC and re-evaluates campaign eligibility."""
        norm = normalize_username(username)
        if not norm:
            return False

        now_utc = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")
        with self._get_connection() as conn:
            conn.execute("DELETE FROM do_not_contact WHERE normalized_username = ?", (norm,))

            lead = conn.execute("SELECT * FROM outreach_leads WHERE normalized_username = ?", (norm,)).fetchone()
            if lead:
                prev_status = lead["review_status"]
                cls_name = lead["mumbai_classification"]

                # Recalculate eligibility
                if cls_name == "Likely Mumbai":
                    new_elig = ELIGIBILITY_MUMBAI_TARGET
                elif cls_name == "Serves Mumbai":
                    new_elig = ELIGIBILITY_BUSINESS_REVIEW
                elif cls_name == "Possibly Mumbai":
                    new_elig = ELIGIBILITY_LOW_CONF_REVIEW
                elif cls_name == "Likely outside Mumbai":
                    new_elig = ELIGIBILITY_SUPPRESSED_OUTSIDE
                else:
                    new_elig = ELIGIBILITY_PENDING_LOCATION

                conn.execute("""
                    UPDATE outreach_leads SET
                        is_dnc = 0,
                        dnc_reason = '',
                        campaign_eligibility = ?,
                        review_status = ?,
                        last_action = 'Removed from DNC',
                        last_action_at = ?,
                        updated_at = ?
                    WHERE normalized_username = ?
                """, (new_elig, STATUS_PENDING, now_utc, now_utc, norm))

                conn.execute("""
                    INSERT INTO outreach_audit_log (
                        normalized_username, action, previous_status, new_status, notes, operator, created_at
                    ) VALUES (?, 'REMOVE_DNC', ?, ?, 'DNC constraint lifted', ?, ?)
                """, (norm, prev_status, STATUS_PENDING, operator, now_utc))

        logger.info("Removed @%s from Do Not Contact registry", username)
        return True

    # -------------------------------------------------------------------------
    # Review & Draft Actions
    # -------------------------------------------------------------------------
    def update_lead_status(
        self,
        username: str,
        new_status: str,
        notes: str = "",
        operator: str = "Operator",
    ) -> bool:
        """
        Updates review status (e.g. Approved, Rejected, Contacted, Replied, Interested, Not Interested).
        Note: Approving a draft MUST NOT send it.
        """
        norm = normalize_username(username)
        if not norm or new_status not in ALL_REVIEW_STATUSES:
            return False

        now_utc = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")
        with self._get_connection() as conn:
            lead = conn.execute("SELECT * FROM outreach_leads WHERE normalized_username = ?", (norm,)).fetchone()
            if not lead:
                return False

            prev_status = lead["review_status"]
            # If lead is marked DNC, enforce DNC unless explicitly removing
            if lead["is_dnc"] and new_status != STATUS_DNC:
                raise DNCSuppressionError(f"Cannot transition @{username} to '{new_status}': account is in Do Not Contact.")

            conn.execute("""
                UPDATE outreach_leads SET
                    review_status = ?,
                    last_action = ?,
                    last_action_at = ?,
                    updated_at = ?
                WHERE normalized_username = ?
            """, (new_status, f"Status -> {new_status}", now_utc, now_utc, norm))

            conn.execute("""
                INSERT INTO outreach_audit_log (
                    normalized_username, action, previous_status, new_status, notes, operator, created_at
                ) VALUES (?, 'STATUS_UPDATE', ?, ?, ?, ?, ?)
            """, (norm, prev_status, new_status, notes, operator, now_utc))

        logger.info("Updated @%s status: %s -> %s", username, prev_status, new_status)
        return True

    def edit_lead_draft(
        self, username: str, new_draft: str, operator: str = "Operator"
    ) -> bool:
        """Updates the drafted message for a lead with audit logging."""
        norm = normalize_username(username)
        if not norm or not new_draft.strip():
            return False

        now_utc = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")
        with self._get_connection() as conn:
            lead = conn.execute("SELECT review_status FROM outreach_leads WHERE normalized_username = ?", (norm,)).fetchone()
            if not lead:
                return False

            conn.execute("""
                UPDATE outreach_leads SET
                    drafted_message = ?,
                    last_action = 'Draft Edited',
                    last_action_at = ?,
                    updated_at = ?
                WHERE normalized_username = ?
            """, (new_draft.strip(), now_utc, now_utc, norm))

            conn.execute("""
                INSERT INTO outreach_audit_log (
                    normalized_username, action, previous_status, new_status, notes, operator, created_at
                ) VALUES (?, 'EDIT_DRAFT', ?, ?, 'Custom message draft updated', ?, ?)
            """, (norm, lead["review_status"], lead["review_status"], operator, now_utc))

        return True

    # -------------------------------------------------------------------------
    # Queries & Pagination
    # -------------------------------------------------------------------------
    def get_lead(self, username: str) -> Optional[Dict[str, Any]]:
        norm = normalize_username(username)
        with self._get_connection() as conn:
            row = conn.execute("SELECT * FROM outreach_leads WHERE normalized_username = ?", (norm,)).fetchone()
            return dict(row) if row else None

    def get_leads_paginated(
        self,
        limit: int = 50,
        offset: int = 0,
        search: str = "",
        eligibility: Optional[str] = None,
        review_status: Optional[str] = None,
        source_brand: Optional[str] = None,
        sort_by: str = "last_action_at DESC",
    ) -> Tuple[List[Dict[str, Any]], int]:
        """
        Retrieves leads using server-side SQL pagination and filtering.
        Returns: (leads_list, total_filtered_count)
        """
        where_clauses = []
        params = []

        if search.strip():
            term = f"%{search.strip().lower()}%"
            where_clauses.append("(normalized_username LIKE ? OR stated_location LIKE ? OR sample_comment_text LIKE ?)")
            params.extend([term, term, term])

        if eligibility:
            where_clauses.append("campaign_eligibility = ?")
            params.append(eligibility)

        if review_status:
            where_clauses.append("review_status = ?")
            params.append(review_status)

        if source_brand:
            where_clauses.append("source_brands LIKE ?")
            params.append(f"%{source_brand.strip().lower()}%")

        where_sql = f"WHERE {' AND '.join(where_clauses)}" if where_clauses else ""

        allowed_sorts = {
            "last_action_at DESC": "last_action_at DESC",
            "last_action_at ASC": "last_action_at ASC",
            "total_comments DESC": "total_comments DESC",
            "username ASC": "username ASC",
        }
        order_sql = allowed_sorts.get(sort_by, "last_action_at DESC")

        with self._get_connection() as conn:
            count_query = f"SELECT COUNT(*) FROM outreach_leads {where_sql}"
            total_count = conn.execute(count_query, params).fetchone()[0]

            data_query = f"""
                SELECT * FROM outreach_leads
                {where_sql}
                ORDER BY {order_sql}
                LIMIT ? OFFSET ?
            """
            rows = conn.execute(data_query, params + [limit, offset]).fetchall()
            leads = [dict(r) for r in rows]

        return leads, total_count

    def get_lead_counts_summary(self) -> Dict[str, int]:
        """Returns aggregate metrics across campaign eligibility and review statuses."""
        with self._get_connection() as conn:
            tot = conn.execute("SELECT COUNT(*) FROM outreach_leads").fetchone()[0]

            def _cnt(col: str, val: Any) -> int:
                return conn.execute(f"SELECT COUNT(*) FROM outreach_leads WHERE {col} = ?", (val,)).fetchone()[0]

            return {
                "total_leads": tot,
                "eligible_mumbai": _cnt("campaign_eligibility", ELIGIBILITY_MUMBAI_TARGET),
                "review_business": _cnt("campaign_eligibility", ELIGIBILITY_BUSINESS_REVIEW),
                "review_low_conf": _cnt("campaign_eligibility", ELIGIBILITY_LOW_CONF_REVIEW),
                "suppressed_outside": _cnt("campaign_eligibility", ELIGIBILITY_SUPPRESSED_OUTSIDE),
                "suppressed_dnc": _cnt("campaign_eligibility", ELIGIBILITY_SUPPRESSED_DNC),
                "pending_location": _cnt("campaign_eligibility", ELIGIBILITY_PENDING_LOCATION),
                "pending_review": _cnt("review_status", STATUS_PENDING),
                "approved": _cnt("review_status", STATUS_APPROVED),
                "contacted": _cnt("review_status", STATUS_CONTACTED),
                "replied": _cnt("review_status", STATUS_REPLIED),
                "interested": _cnt("review_status", STATUS_INTERESTED),
                "not_interested": _cnt("review_status", STATUS_NOT_INTERESTED),
                "rejected": _cnt("review_status", STATUS_REJECTED),
            }

    def get_audit_log(self, limit: int = 50) -> List[Dict[str, Any]]:
        with self._get_connection() as conn:
            rows = conn.execute(
                "SELECT * FROM outreach_audit_log ORDER BY id DESC LIMIT ?", (limit,)
            ).fetchall()
            return [dict(r) for r in rows]

    def get_all_leads_for_export(self) -> List[Dict[str, Any]]:
        with self._get_connection() as conn:
            rows = conn.execute("SELECT * FROM outreach_leads ORDER BY updated_at DESC").fetchall()
            return [dict(r) for r in rows]


# -----------------------------------------------------------------------------
# Official Meta Authorized Messaging Engine
# -----------------------------------------------------------------------------
class MetaMessagingClient:
    """
    Official Meta Messenger API for Instagram messaging client.

    Compliance Rules:
    1. Direct unsolicited messaging to arbitrary Instagram handles is prohibited
       by Meta's Graph API. Messaging requires an existing conversation (24h window)
       and an Instagram Scoped User ID (IGSID).
    2. Accounts without active authorized API channels are routed to
       'Manual / Opt-In Required' with compliant direct links (https://ig.me/m/{username}).
    3. Strictly blocks duplicate sends, DNC violations, and outside-Mumbai exclusions.
    4. NEVER implements automated follower bots or unapproved scraping endpoints.
    """

    def __init__(
        self,
        page_access_token: Optional[str] = None,
        page_id: Optional[str] = None,
        db: Optional[LeadOutreachDB] = None,
    ):
        self.page_access_token = page_access_token or os.environ.get("META_PAGE_ACCESS_TOKEN", "")
        self.page_id = page_id or os.environ.get("META_PAGE_ID", "")
        self.db = db or LeadOutreachDB()

    def validate_outreach_attempt(
        self, lead: Dict[str, Any], campaign: str = CAMPAIGN_MUMBAI
    ) -> None:
        """Validates that a lead satisfies all safety, DNC, and eligibility rules before sending."""
        username = lead.get("username", "")

        # 1. Do Not Contact check
        if lead.get("is_dnc") or lead.get("review_status") == STATUS_DNC:
            raise DNCSuppressionError(f"Outreach blocked: @{username} is on the Do Not Contact list.")

        # 2. Mumbai campaign suppression
        if campaign == CAMPAIGN_MUMBAI and lead.get("campaign_eligibility") == ELIGIBILITY_SUPPRESSED_OUTSIDE:
            raise CampaignSuppressionError(
                f"Outreach blocked: @{username} is confirmed outside Mumbai and suppressed from Mumbai campaign."
            )

        # 3. Duplicate-send protection
        if lead.get("review_status") == STATUS_CONTACTED:
            raise DuplicateSendBlockedError(
                f"Duplicate send blocked: @{username} has already been marked Contacted."
            )

        # 4. Approval check: draft must be approved
        if lead.get("review_status") != STATUS_APPROVED:
            raise UnapprovedDraftError(
                f"Outreach blocked: Draft for @{username} must be in '{STATUS_APPROVED}' status before sending."
            )

    def execute_messaging_workflow(
        self,
        username: str,
        igsid: Optional[str] = None,
        mock: bool = False,
        operator: str = "Operator",
    ) -> Dict[str, Any]:
        """
        Executes compliant messaging workflow for a lead.
        If mock=True: Simulates successful authorized transmission.
        If real: Evaluates credentials and IGSID. If not available, returns
        'Manual / Opt-In Required' status with direct ig.me link.
        """
        lead = self.db.get_lead(username)
        if not lead:
            raise LeadOutreachError(f"Lead @{username} not found in database.")

        # Validate constraints
        self.validate_outreach_attempt(lead, campaign=lead.get("campaign", CAMPAIGN_MUMBAI))

        now_utc = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")

        # Mock Mode for Automated Testing
        if mock:
            self.db.update_lead_status(
                username=username,
                new_status=STATUS_CONTACTED,
                notes=f"Message sent via authorized API (Mock Mode). Recipient IGSID: {igsid or 'MOCK_IGSID_123'}",
                operator=operator,
            )
            with self.db._get_connection() as conn:
                conn.execute(
                    "UPDATE outreach_leads SET messaging_status = ?, last_action_at = ? WHERE normalized_username = ?",
                    (MSG_STATUS_SENT, now_utc, normalize_username(username)),
                )
            return {
                "status": "success",
                "messaging_status": MSG_STATUS_SENT,
                "review_status": STATUS_CONTACTED,
                "recipient": username,
                "message": lead["drafted_message"],
                "channel": "meta_graph_api_mock",
            }

        # Real Mode API Inspection
        if not self.page_access_token or not self.page_id or not igsid:
            # Official Meta API cannot message arbitrary public handles without an IGSID
            # Mark lead as 'Manual / Opt-In Required'
            with self.db._get_connection() as conn:
                conn.execute(
                    "UPDATE outreach_leads SET messaging_status = ?, last_action = 'Manual Outreach Required', last_action_at = ? WHERE normalized_username = ?",
                    (MSG_STATUS_MANUAL, now_utc, normalize_username(username)),
                )
            return {
                "status": "manual_required",
                "messaging_status": MSG_STATUS_MANUAL,
                "review_status": lead["review_status"],
                "recipient": username,
                "reason": "Cold commenter lacks Meta IGSID. Direct message URL provided for manual operator outreach.",
                "direct_message_url": lead["direct_message_url"],
                "drafted_message": lead["drafted_message"],
            }

        # Real Meta Graph API Call (when credentials and IGSID exist)
        try:
            import urllib.request
            url = f"https://graph.facebook.com/v19.0/{self.page_id}/messages"
            payload = {
                "recipient": {"id": igsid},
                "message": {"text": lead["drafted_message"]},
                "access_token": self.page_access_token,
            }
            req = urllib.request.Request(
                url,
                data=json.dumps(payload).encode("utf-8"),
                headers={"Content-Type": "application/json"},
                method="POST",
            )
            with urllib.request.urlopen(req, timeout=10.0) as resp:
                data = json.loads(resp.read().decode("utf-8"))

            self.db.update_lead_status(
                username=username,
                new_status=STATUS_CONTACTED,
                notes=f"Message sent via Meta Graph API. Message ID: {data.get('message_id')}",
                operator=operator,
            )
            with self.db._get_connection() as conn:
                conn.execute(
                    "UPDATE outreach_leads SET messaging_status = ?, last_action_at = ? WHERE normalized_username = ?",
                    (MSG_STATUS_SENT, now_utc, normalize_username(username)),
                )
            return {
                "status": "success",
                "messaging_status": MSG_STATUS_SENT,
                "review_status": STATUS_CONTACTED,
                "response": data,
            }
        except Exception as e:
            logger.error("Meta Graph API error contacting @%s: %s", username, e)
            with self.db._get_connection() as conn:
                conn.execute(
                    "UPDATE outreach_leads SET messaging_status = ?, last_action_at = ? WHERE normalized_username = ?",
                    (MSG_STATUS_FAILED, now_utc, normalize_username(username)),
                )
            raise MetaAuthorizationError(f"Meta Graph API error: {e}") from e


# -----------------------------------------------------------------------------
# Dual Reporting Synchronization (Google Sheets & Excel)
# -----------------------------------------------------------------------------
class LeadOutreachSync:
    """
    Synchronizes outreach leads to:
    1. Google Spreadsheet master tab: 'Outreach Manager'
    2. Local Excel mirror tab: 'Outreach Manager' in exports/instagram_research.xlsx
    Preserves all existing worksheets and never produces CSVs.
    """

    SHEET_TITLE = "Outreach Manager"

    COLOR_MAP = {
        ELIGIBILITY_MUMBAI_TARGET: {"bg_hex": "D1E7DD", "font_hex": "0F5132", "rgb": (0.82, 0.91, 0.87)},
        ELIGIBILITY_BUSINESS_REVIEW: {"bg_hex": "CFF4FC", "font_hex": "055160", "rgb": (0.81, 0.96, 0.99)},
        ELIGIBILITY_LOW_CONF_REVIEW: {"bg_hex": "FFF3CD", "font_hex": "664D03", "rgb": (1.0, 0.95, 0.80)},
        ELIGIBILITY_SUPPRESSED_OUTSIDE: {"bg_hex": "E2E3E5", "font_hex": "41464B", "rgb": (0.89, 0.89, 0.90)},
        ELIGIBILITY_SUPPRESSED_DNC: {"bg_hex": "F8D7DA", "font_hex": "842029", "rgb": (0.97, 0.84, 0.85)},
        ELIGIBILITY_PENDING_LOCATION: {"bg_hex": "F8D7DA", "font_hex": "842029", "rgb": (0.97, 0.84, 0.85)},
    }

    @classmethod
    def sync_to_excel(
        cls,
        leads: Optional[List[Dict[str, Any]]] = None,
        excel_path: str = DEFAULT_EXCEL_PATH,
        db: Optional[LeadOutreachDB] = None,
    ) -> int:
        """
        Creates or updates 'Outreach Manager' tab in exports/instagram_research.xlsx,
        preserving all other worksheets.
        """
        import openpyxl
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter

        target = Path(excel_path)
        target.parent.mkdir(parents=True, exist_ok=True)

        if target.exists():
            wb = openpyxl.load_workbook(str(target))
        else:
            wb = openpyxl.Workbook()
            if "Sheet" in wb.sheetnames:
                wb.remove(wb["Sheet"])

        if leads is None:
            db_inst = db or LeadOutreachDB()
            leads = db_inst.get_all_leads_for_export()

        if cls.SHEET_TITLE in wb.sheetnames:
            ws = wb[cls.SHEET_TITLE]
            ws.delete_rows(1, ws.max_row + 1)
        else:
            ws = wb.create_sheet(title=cls.SHEET_TITLE)

        now_str = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")

        # Formatting Styles
        font_title = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        fill_title = PatternFill("solid", fgColor="1F4E78")
        font_hdr = Font(name="Calibri", size=9.5, bold=True, color="FFFFFF")
        fill_hdr = PatternFill("solid", fgColor="2F343D")
        font_body = Font(name="Calibri", size=9.5, color="1B2631")
        font_link = Font(name="Calibri", size=9.5, bold=True, underline="single", color="0563C1")
        thin_border = Border(
            left=Side(style="thin", color="D5D8DC"),
            right=Side(style="thin", color="D5D8DC"),
            top=Side(style="thin", color="D5D8DC"),
            bottom=Side(style="thin", color="D5D8DC"),
        )

        # Banner Row
        ws.row_dimensions[1].height = 26
        banner_vals = [
            (1, "MEATNDOOR LEAD OUTREACH MANAGER"),
            (2, f"Last Updated: {now_str}"),
            (3, f"Total Leads in Queue: {len(leads)}"),
        ]
        for col, txt in banner_vals:
            c = ws.cell(row=1, column=col, value=txt)
            c.font = font_title
            c.fill = fill_title
            c.alignment = Alignment(vertical="center", horizontal="left" if col <= 2 else "center")

        ws.row_dimensions[2].height = 10

        # Table Headers
        headers = [
            "Username", "Profile URL", "Direct Chat (ig.me)", "Source Brands",
            "Source Posts", "Total Comments", "Location Classification", "Confidence",
            "Stated Location", "Campaign Eligibility", "Review Status",
            "Drafted Outreach Message", "Messaging Status", "Last Action", "Updated UTC"
        ]
        ws.row_dimensions[3].height = 22
        for c_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=3, column=c_idx, value=h)
            cell.font = font_hdr
            cell.fill = fill_hdr
            cell.alignment = Alignment(vertical="center", horizontal="center" if c_idx in (6, 8, 10, 11, 13) else "left")
            cell.border = thin_border

        # Populate Rows
        for r_idx, lead in enumerate(leads, start=4):
            ws.row_dimensions[r_idx].height = 20
            u = lead.get("username", "")
            p_url = lead.get("profile_url", f"https://www.instagram.com/{u}/")
            dm_url = lead.get("direct_message_url", f"https://ig.me/m/{u}")
            brands = lead.get("source_brands", "[]")
            if brands.startswith("["):
                try:
                    brands = ", ".join(json.loads(brands))
                except Exception:
                    pass
            posts = lead.get("source_posts", "[]")
            if posts.startswith("["):
                try:
                    posts = ", ".join(json.loads(posts)[:5])
                except Exception:
                    pass

            row_data = [
                f"@{u}",
                p_url,
                dm_url,
                brands,
                posts,
                lead.get("total_comments", 1),
                lead.get("mumbai_classification", "Not enough evidence"),
                lead.get("mumbai_confidence", "Low"),
                lead.get("stated_location", "—"),
                lead.get("campaign_eligibility", ELIGIBILITY_PENDING_LOCATION),
                lead.get("review_status", STATUS_PENDING),
                lead.get("drafted_message", ""),
                lead.get("messaging_status", MSG_STATUS_MANUAL),
                lead.get("last_action", "Ingested"),
                lead.get("updated_at", now_str),
            ]

            for c_idx, val in enumerate(row_data, start=1):
                cell = ws.cell(row=r_idx, column=c_idx, value=val)
                cell.font = font_body
                cell.border = thin_border
                cell.alignment = Alignment(
                    vertical="center",
                    horizontal="center" if c_idx in (1, 6, 8, 10, 11, 13, 15) else "left"
                )

                if c_idx in (2, 3):
                    cell.hyperlink = val
                    cell.font = font_link

            # Color badge for Campaign Eligibility
            elig_cell = ws.cell(row=r_idx, column=10)
            elig_val = lead.get("campaign_eligibility", ELIGIBILITY_PENDING_LOCATION)
            cfg = cls.COLOR_MAP.get(elig_val, cls.COLOR_MAP[ELIGIBILITY_PENDING_LOCATION])
            elig_cell.fill = PatternFill("solid", fgColor=cfg["bg_hex"])
            elig_cell.font = Font(name="Calibri", size=9.5, bold=True, color=cfg["font_hex"])

        # Auto-fit column widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                if cell.row == 1:
                    continue
                max_len = max(max_len, len(val))
            ws.column_dimensions[col_letter].width = max(14, min(max_len + 3, 50))

        wb.save(str(target))
        logger.info("Successfully updated '%s' worksheet in %s (%d rows)", cls.SHEET_TITLE, target, len(leads))
        return len(leads)

    @classmethod
    def sync_to_google_sheets(
        cls,
        leads: Optional[List[Dict[str, Any]]] = None,
        spreadsheet_id: str = DEFAULT_SPREADSHEET_ID,
        credentials_path: Optional[str] = None,
        db: Optional[LeadOutreachDB] = None,
    ) -> Dict[str, Any]:
        """
        Creates or updates 'Outreach Manager' tab in the master Google Spreadsheet,
        preserving all other worksheets.
        """
        from sheets_sync import GoogleSheetsManager

        if leads is None:
            db_inst = db or LeadOutreachDB()
            leads = db_inst.get_all_leads_for_export()

        mgr = GoogleSheetsManager(spreadsheet_id=spreadsheet_id)
        service = mgr.get_service(credentials_path=credentials_path)

        meta = service.spreadsheets().get(spreadsheetId=spreadsheet_id).execute()
        sheet_map = {s["properties"]["title"]: s["properties"]["sheetId"] for s in meta.get("sheets", [])}

        sheet_title = cls.SHEET_TITLE
        if sheet_title not in sheet_map:
            add_req = {
                "requests": [
                    {
                        "addSheet": {
                            "properties": {
                                "title": sheet_title,
                                "gridProperties": {"rowCount": max(100, len(leads) + 15), "columnCount": 16},
                            }
                        }
                    }
                ]
            }
            res = service.spreadsheets().batchUpdate(spreadsheetId=spreadsheet_id, body=add_req).execute()
            sheet_id = res["replies"][0]["addSheet"]["properties"]["sheetId"]
        else:
            sheet_id = sheet_map[sheet_title]

        now_str = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")

        headers = [
            "Username", "Profile URL", "Direct Chat (ig.me)", "Source Brands",
            "Source Posts", "Total Comments", "Location Classification", "Confidence",
            "Stated Location", "Campaign Eligibility", "Review Status",
            "Drafted Outreach Message", "Messaging Status", "Last Action", "Updated UTC"
        ]

        # Build values grid
        values = [
            ["MEATNDOOR LEAD OUTREACH MANAGER", f"Last Updated: {now_str}", f"Total Leads in Queue: {len(leads)}"],
            [],
            headers,
        ]

        for lead in leads:
            u = lead.get("username", "")
            p_url = lead.get("profile_url", f"https://www.instagram.com/{u}/")
            dm_url = lead.get("direct_message_url", f"https://ig.me/m/{u}")
            brands = lead.get("source_brands", "[]")
            if brands.startswith("["):
                try:
                    brands = ", ".join(json.loads(brands))
                except Exception:
                    pass
            posts = lead.get("source_posts", "[]")
            if posts.startswith("["):
                try:
                    posts = ", ".join(json.loads(posts)[:5])
                except Exception:
                    pass

            values.append([
                f'=HYPERLINK("{p_url}", "@{u}")',
                p_url,
                f'=HYPERLINK("{dm_url}", "Direct Chat")',
                brands,
                posts,
                lead.get("total_comments", 1),
                lead.get("mumbai_classification", "Not enough evidence"),
                lead.get("mumbai_confidence", "Low"),
                lead.get("stated_location", "—"),
                lead.get("campaign_eligibility", ELIGIBILITY_PENDING_LOCATION),
                lead.get("review_status", STATUS_PENDING),
                lead.get("drafted_message", ""),
                lead.get("messaging_status", MSG_STATUS_MANUAL),
                lead.get("last_action", "Ingested"),
                lead.get("updated_at", now_str),
            ])

        # Clear and update
        service.spreadsheets().values().clear(
            spreadsheetId=spreadsheet_id, range=f"'{sheet_title}'!A1:Z{max(200, len(values) + 20)}"
        ).execute()

        service.spreadsheets().values().update(
            spreadsheetId=spreadsheet_id,
            range=f"'{sheet_title}'!A1",
            valueInputOption="USER_ENTERED",
            body={"values": values},
        ).execute()

        # Apply formatting requests (header styling, freeze panes)
        fmt_requests = [
            # Freeze row 3
            {
                "updateSheetProperties": {
                    "properties": {
                        "sheetId": sheet_id,
                        "gridProperties": {"frozenRowCount": 3},
                    },
                    "fields": "gridProperties.frozenRowCount",
                }
            },
            # Style Header Banner (Row 1)
            {
                "repeatCell": {
                    "range": {"sheetId": sheet_id, "startRowIndex": 0, "endRowIndex": 1, "startColumnIndex": 0, "endColumnIndex": 15},
                    "cell": {
                        "userEnteredFormat": {
                            "backgroundColor": {"red": 0.12, "green": 0.31, "blue": 0.47},
                            "textFormat": {"bold": True, "foregroundColor": {"red": 1.0, "green": 1.0, "blue": 1.0}, "fontSize": 11},
                            "verticalAlignment": "MIDDLE",
                        }
                    },
                    "fields": "userEnteredFormat(backgroundColor,textFormat,verticalAlignment)",
                }
            },
            # Style Table Headers (Row 3)
            {
                "repeatCell": {
                    "range": {"sheetId": sheet_id, "startRowIndex": 2, "endRowIndex": 3, "startColumnIndex": 0, "endColumnIndex": 15},
                    "cell": {
                        "userEnteredFormat": {
                            "backgroundColor": {"red": 0.18, "green": 0.20, "blue": 0.24},
                            "textFormat": {"bold": True, "foregroundColor": {"red": 1.0, "green": 1.0, "blue": 1.0}, "fontSize": 10},
                            "verticalAlignment": "MIDDLE",

                        }
                    },
                    "fields": "userEnteredFormat(backgroundColor,textFormat,verticalAlignment)",
                }
            },
        ]

        service.spreadsheets().batchUpdate(
            spreadsheetId=spreadsheet_id, body={"requests": fmt_requests}
        ).execute()

        return {
            "spreadsheet_id": spreadsheet_id,
            "sheet_title": sheet_title,
            "rows_synced": len(leads),
            "spreadsheet_url": f"https://docs.google.com/spreadsheets/d/{spreadsheet_id}/edit#gid={sheet_id}",
        }


# -----------------------------------------------------------------------------
# CLI Entrypoint
# -----------------------------------------------------------------------------
def main():
    import argparse
    parser = argparse.ArgumentParser(description="MeatNDoor Lead Outreach Manager CLI")
    parser.add_argument("--ingest", action="store_true", help="Ingest leads from existing checkpoint DBs")
    parser.add_argument("--source", nargs="+", help="Specific brand accounts to ingest from")
    parser.add_argument("--sync-excel", action="store_true", help="Sync outreach leads to exports/instagram_research.xlsx")
    parser.add_argument("--sync-sheets", action="store_true", help="Sync outreach leads to Google Sheets")
    parser.add_argument("--dnc-add", help="Add username to Do Not Contact list")
    parser.add_argument("--dnc-reason", default="Manual CLI suppression", help="Reason for DNC")
    parser.add_argument("--dnc-remove", help="Remove username from Do Not Contact list")
    args = parser.parse_args()

    db = LeadOutreachDB()

    if args.dnc_add:
        db.add_to_dnc(args.dnc_add, reason=args.dnc_reason, operator="CLI")
        print(f"[✓] Added @{args.dnc_add} to Do Not Contact.")

    if args.dnc_remove:
        db.remove_from_dnc(args.dnc_remove, operator="CLI")
        print(f"[✓] Removed @{args.dnc_remove} from Do Not Contact.")

    if args.ingest:
        print("[*] Ingesting leads from checkpoints...")
        stats = db.ingest_from_checkpoints(source_accounts=args.source)
        print(f"[✓] Ingestion completed:")
        print(f"    - Raw comments scanned: {stats['total_raw_found']}")
        print(f"    - Unique leads processed: {stats['leads_ingested_or_updated']}")
        print(f"    - New leads added: {stats['new_leads_added']}")
        print(f"    - Mumbai eligible leads: {stats['mumbai_eligible']}")
        print(f"    - DNC suppressed leads: {stats['dnc_suppressed']}")

    if args.sync_excel:
        print("[*] Syncing outreach queue to local Excel mirror...")
        cnt = LeadOutreachSync.sync_to_excel(db=db)
        print(f"[✓] Synced {cnt} leads to {DEFAULT_EXCEL_PATH} -> 'Outreach Manager'.")

    if args.sync_sheets:
        print("[*] Syncing outreach queue to Google Spreadsheet...")
        try:
            res = LeadOutreachSync.sync_to_google_sheets(db=db)
            print(f"[✓] Synced {res['rows_synced']} leads to Google Sheet: {res['spreadsheet_url']}")
        except Exception as e:
            print(f"[!] Google Sheets sync error: {e}")


if __name__ == "__main__":
    main()
